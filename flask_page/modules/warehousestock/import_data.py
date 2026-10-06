import os
import openpyxl
from datetime import datetime
from config import Config
from app import create_app
from app.models import db, User, Customer, Item, ItemStock, ProductionPlan, DailyStockRecord, DeliverySchedule, StockMovement

app = create_app()

def parse_float(val, default=0.0):
    if val is None or val == '' or str(val).strip() in ['#REF!', '-', 'None']:
        return default
    try:
        return float(val)
    except (ValueError, TypeError):
        return default

def clean_str(val):
    if val is None:
        return ''
    s = str(val).strip()
    return '' if s == 'None' else s

def import_all():
    excel_path = Config.EXCEL_SOURCE_PATH
    if not os.path.exists(excel_path):
        print(f"Excel file not found at: {excel_path}")
        return

    print(f"Loading workbook: {excel_path} (read_only mode)...")
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)

    with app.app_context():
        # Create tables
        db.create_all()

        # 1. Seed Users
        print("\n--- Seeding Users ---")
        users_to_seed = [
            ('admin', 'admin@warehouse.local', 'adminpassword123', 'Admin', 'Chief Admin'),
            ('manager', 'manager@warehouse.local', 'manager123', 'Manager', 'Warehouse Manager'),
            ('operator', 'operator@warehouse.local', 'operator123', 'Operator', 'Stock Operator')
        ]
        for username, email, pwd, role, fullname in users_to_seed:
            u = User.query.filter_by(username=username).first()
            if not u:
                u = User(username=username, email=email, role=role, full_name=fullname)
                u.set_password(pwd)
                db.session.add(u)
                print(f"  Created user: {username} (Password: {pwd}) [{role}]")
            else:
                print(f"  User {username} already exists.")
        db.session.commit()

        # 2. Seed Customers
        print("\n--- Seeding Customers ---")
        customer_names = [
            'PIDMY',
            'Nichicon(M)',
            'Nichicon(Wuxi)',
            'Panasonic(I)',
            'Panasonic Saga',
            'Beading type  - From Japan'
        ]
        customer_map = {}
        for cname in customer_names:
            c = Customer.query.filter_by(name=cname).first()
            if not c:
                c = Customer(name=cname, code=cname[:6].upper().replace(' ', ''))
                db.session.add(c)
                db.session.flush()
                print(f"  Added Customer: {cname}")
            customer_map[cname] = c
        db.session.commit()

        # 3. Import Items & MC Plan
        print("\n--- Importing Items and MC Plan ---")
        if 'mc plan' in wb.sheetnames:
            mc_sheet = wb['mc plan']
            current_customer = customer_map['PIDMY']
            items_added = 0

            for row in mc_sheet.iter_rows(values_only=True):
                col0 = clean_str(row[0])
                if not col0 or col0 in ['Stock as at', 'Items', 'r/i']:
                    continue

                # Check if col0 is a customer header
                if col0 in customer_map and row[1] is None:
                    current_customer = customer_map[col0]
                    continue

                # It's an item row if row[1] is present
                if row[1] is not None:
                    part_no = col0
                    std_pkg = parse_float(row[1], default=1.0)
                    mc_num = clean_str(row[3])
                    mat_size = clean_str(row[4])
                    stock_kpc = parse_float(row[5])
                    stock_ctn = parse_float(row[6])
                    closing_ctn = parse_float(row[8])
                    closing_kpcs = parse_float(row[9])
                    m1_forecast = parse_float(row[10])
                    m2_forecast = parse_float(row[11])
                    shortage_kpcs = parse_float(row[12])
                    shortage_ctn = parse_float(row[13])

                    item = Item.query.filter_by(part_no=part_no).first()
                    if not item:
                        item = Item(
                            part_no=part_no,
                            customer_id=current_customer.id,
                            std_pkg=std_pkg,
                            mc_number=mc_num,
                            material_size=mat_size
                        )
                        db.session.add(item)
                        db.session.flush()
                        items_added += 1

                        # Initial stock
                        stock = ItemStock(
                            item_id=item.id,
                            warehouse_ctn=closing_ctn,
                            finished_goods_ctn=stock_ctn,
                            total_ctn=closing_ctn,
                            total_kpcs=closing_kpcs
                        )
                        db.session.add(stock)

                        # Production plan
                        plan = ProductionPlan(
                            item_id=item.id,
                            closing_ctn=closing_ctn,
                            closing_kpcs=closing_kpcs,
                            forecast_month1=m1_forecast,
                            forecast_month2=m2_forecast,
                            shortage_kpcs=shortage_kpcs,
                            shortage_ctn=shortage_ctn
                        )
                        plan.calculate_shortage(std_pkg)
                        db.session.add(plan)

            db.session.commit()
            print(f"  Imported {items_added} items from 'mc plan'.")

        # 4. Import Stock Control Channels
        print("\n--- Importing Stock Control Channel stages ---")
        if 'stock control' in wb.sheetnames:
            sc_sheet = wb['stock control']
            sc_updated = 0
            for row in sc_sheet.iter_rows(values_only=True):
                part_no = clean_str(row[1]) if len(row) > 1 else ''
                channel = clean_str(row[2]) if len(row) > 2 else ''
                if not part_no or not channel:
                    continue

                item = Item.query.filter_by(part_no=part_no).first()
                if item and item.stock:
                    # Find latest value in row (e.g. from column 3 onwards)
                    latest_val = 0.0
                    for val in reversed(row[3:35]):
                        if val is not None and val != '' and val != '#REF!':
                            latest_val = parse_float(val)
                            break

                    ch_lower = channel.lower()
                    if 'warehouse' in ch_lower:
                        item.stock.warehouse_ctn = latest_val
                    elif 'qa transit' in ch_lower:
                        item.stock.qa_transit_ctn = latest_val
                    elif 'return qa' in ch_lower:
                        item.stock.return_qa_ctn = latest_val
                    elif 'wash' in ch_lower:
                        item.stock.wash_area_ctn = latest_val
                    elif 'qa area' in ch_lower:
                        item.stock.qa_area_ctn = latest_val
                    elif 'production' in ch_lower:
                        item.stock.production_ctn = latest_val
                    elif 'finished' in ch_lower:
                        item.stock.finished_goods_ctn = latest_val
                    elif 'ng' in ch_lower or 'disposed' in ch_lower:
                        item.stock.ng_disposed_ctn = latest_val
                    
                    item.stock.recalculate(item.std_pkg)
                    sc_updated += 1
            
            db.session.commit()
            print(f"  Updated channel balances for {sc_updated} stage records.")

        # 5. Import Daily Stock Sheets (29.04.15 to 09.05.15)
        print("\n--- Importing Daily Stock & Delivery Schedules ---")
        daily_sheets = [s for s in wb.sheetnames if s not in ['stock control', 'master', 'mc plan', 'washing qa']]
        total_daily = 0
        total_deliv = 0

        for sname in daily_sheets:
            sheet = wb[sname]
            sheet_rows = list(sheet.iter_rows(values_only=True))
            if len(sheet_rows) < 5:
                continue

            for row in sheet_rows[4:]:
                part_no = clean_str(row[0])
                if not part_no or part_no in ['PART NO.', 'Items', 'Total']:
                    continue

                item = Item.query.filter_by(part_no=part_no).first()
                if not item:
                    # If item not yet created, create it under PIDMY
                    item = Item(
                        part_no=part_no,
                        customer_id=customer_map['PIDMY'].id,
                        std_pkg=parse_float(row[5], default=1.0)
                    )
                    db.session.add(item)
                    db.session.flush()
                    stock = ItemStock(item_id=item.id, warehouse_ctn=parse_float(row[1]))
                    stock.recalculate(item.std_pkg)
                    db.session.add(stock)

                pedma_part = clean_str(row[11]) if len(row) > 11 else ''
                if pedma_part and not item.delivery_part_no:
                    item.delivery_part_no = pedma_part

                wh_stock = parse_float(row[1]) if len(row) > 1 else 0.0
                in_qty = parse_float(row[2]) if len(row) > 2 else 0.0
                out_qty = parse_float(row[3]) if len(row) > 3 else 0.0
                qty_ctn = parse_float(row[4]) if len(row) > 4 else 0.0
                heating_ctn = parse_float(row[6]) if len(row) > 6 else 0.0
                qa_ctn = parse_float(row[7]) if len(row) > 7 else 0.0
                wh_ctn = parse_float(row[8]) if len(row) > 8 else 0.0
                kpcs = parse_float(row[9]) if len(row) > 9 else 0.0

                daily_rec = DailyStockRecord(
                    item_id=item.id,
                    record_date=sname,
                    wh_stock=wh_stock,
                    in_qty=in_qty,
                    out_qty=out_qty,
                    qty_per_ctn=qty_ctn,
                    heating_ctn=heating_ctn,
                    qa_ctn=qa_ctn,
                    wh_ctn=wh_ctn,
                    total_ctn=wh_stock,
                    kpcs=kpcs
                )
                db.session.add(daily_rec)
                total_daily += 1

                # Delivery schedule slots if present (Pedma delivery 1, 2, 3...)
                if len(row) > 14 and pedma_part:
                    stk1 = parse_float(row[12])
                    qty1 = parse_float(row[13])
                    bal1 = parse_float(row[14])
                    if stk1 > 0 or qty1 > 0 or bal1 > 0:
                        deliv = DeliverySchedule(
                            item_id=item.id,
                            delivery_date=sname,
                            delivery_label="Delivery 1",
                            pedma_part_no=pedma_part,
                            stock_qty=stk1,
                            deliver_qty=qty1,
                            balance_qty=bal1
                        )
                        db.session.add(deliv)
                        total_deliv += 1

            db.session.commit()
            print(f"  Processed daily sheet: {sname}")

        print(f"\nImport finished successfully!")
        print(f"Total Daily Stock Records: {total_daily}")
        print(f"Total Delivery Schedules: {total_deliv}")

    wb.close()

if __name__ == '__main__':
    import_all()
